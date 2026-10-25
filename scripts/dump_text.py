"""Writes plain-text dumps of the three deliverables for review (no images).

Run from the repository root:  python3 scripts/dump_text.py <output_dir>
Produces notebook.txt (markdown cells, code and text outputs), report.md (via pandoc) and slides.md (via python-pptx).
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)

# --- notebook ---
nb = json.loads((ROOT / "notebooks" / "Loan_Default_Prediction_Capstone.ipynb").read_text())
lines = []
for i, c in enumerate(nb["cells"]):
    src = "".join(c["source"])
    if c["cell_type"] == "markdown":
        lines.append(f"\n===== CELL {i} [markdown] =====\n{src}\n")
        continue
    lines.append(f"\n===== CELL {i} [code] =====\n{src}\n----- output -----")
    for o in c.get("outputs", []):
        t = o.get("output_type")
        if t == "stream":
            lines.append("".join(o["text"]))
        elif t in ("execute_result", "display_data"):
            d = o.get("data", {})
            if "text/html" in d and "text/plain" in d and "Styler" in "".join(d["text/plain"]):
                h = "".join(d["text/html"])
                h = re.sub(r"<style.*?</style>", "", h, flags=re.S)
                for r in re.findall(r"<tr>(.*?)</tr>", h, flags=re.S):
                    cells = re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, flags=re.S)
                    lines.append(" | ".join(re.sub("<.*?>", "", x).strip() for x in cells))
            elif "text/plain" in d:
                tp = "".join(d["text/plain"])
                lines.append("[figure]" if "Figure" in tp and "Axes" in tp else tp)
        elif t == "error":
            lines.append(f"ERROR: {o.get('ename')}: {o.get('evalue')}")
(OUT / "notebook.txt").write_text("\n".join(lines))

# --- report ---
subprocess.run(["pandoc", str(ROOT / "reports" / "Loan_Default_Prediction_Final_Report.docx"), "-t", "markdown",
                "-o", str(OUT / "report.md")], check=True)

# --- slides ---
prs = Presentation(str(ROOT / "presentation" / "Loan_Default_Prediction_Presentation.pptx"))
out = []
for n, slide in enumerate(prs.slides, 1):
    out.append(f"\n## Slide {n}")
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            out.append(shape.text_frame.text.strip())
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                out.append(" | ".join(cell.text for cell in row.cells))
        if getattr(shape, "has_chart", False) and shape.has_chart:
            ch = shape.chart
            title = ch.chart_title.text_frame.text if ch.has_title else "(chart)"
            out.append(f"[chart: {title}]")
            cats = list(ch.plots[0].categories)
            for series in ch.plots[0].series:
                out.append(f"  series {series.name}: " + ", ".join(f"{c}={v}" for c, v in zip(cats, series.values)))
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
        out.append("[speaker notes] " + slide.notes_slide.notes_text_frame.text.strip())
(OUT / "slides.md").write_text("\n".join(out))
for f in ("notebook.txt", "report.md", "slides.md"):
    print(f, (OUT / f).stat().st_size, "bytes")
